#импорт библиотек и нужных модулей
import os

from flask import Blueprint, jsonify, request, url_for, current_app
from flask_login import login_required, current_user, login_user

from quiz_app import db
from quiz_app.models import User, Event, Team, RegistrationsEvent, SitePoster
from datetime import datetime

from ..helpers import _format_short_name, _require_admin_json, _MONTHS_RU, _format_event, _save_upload, _format_date_ru, _save_image_upload




api_admin_bp = Blueprint("api_admin", __name__, url_prefix="/api")



#Загрузка фото для мероприятия
@api_admin_bp.route("/upload/event-photo", methods=["POST"])
@login_required
def upload_event_photo():
    
    if not _require_admin_json():
        return jsonify({"status": "error", "message": "Нет доступа"}), 403
    
    if "file" not in request.files:
        return jsonify({"status": "error", "message": "Файл не передан"}), 400
    
    file = request.files["file"]
    rel_path = _save_upload(file, "events", max_size=(2000, 2000))
    
    if not rel_path:
        return jsonify({"status": "error", "message": "Недопустимый формат файла"}), 400
    photo_url = url_for("pages.serve_media", filename=rel_path)
    
    return jsonify({"status": "success", "url": photo_url, "path": rel_path})



#Загрузка фото для результатов
@api_admin_bp.route("/upload/result-photo", methods=["POST"])
@login_required
def upload_result_photo():
    
    if not _require_admin_json():
        return jsonify({"status": "error", "message": "No access"}), 403
    
    if "file" not in request.files:
        return jsonify({"status": "error", "message": "File is required"}), 400
    
    file = request.files["file"]
    
    #Сохранение фото в папке results
    rel_path = _save_image_upload(
        file,
        "results",
        max_file_size=16 * 1024 * 1024,
    )
    
    if not rel_path:
        return jsonify({"status": "error", "message": "Недопустимый формат файла или файл больше 16 МБ"}), 400
    
    photo_url = url_for("pages.serve_media", filename=rel_path)
    
    #Возвращаем URL фото и путь к файлу
    return jsonify({"status": "success", "url": photo_url, "path": rel_path})


#Недавние игры в блоке Недавние игры
@api_admin_bp.route("/admin/recent-games")
def admin_recent_games():
    if not _require_admin_json():
        return jsonify({"status": "error", "message": "Нет доступа"}), 403
    from datetime import datetime, timedelta
    now = datetime.now()
    month_start = datetime(now.year, now.month, 1)
    month_end = (month_start + timedelta(days=32)).replace(day=1) - timedelta(seconds=1)
    events = Event.query.filter(
        Event.date >= month_start,
        Event.date <= month_end,
    ).order_by(Event.date.asc()).all()
    rows = []
    for e in events:
        d = e.date
        rows.append(
            {
                "id": e.id,
                "title": e.name,
                "date": f"{d.day} {_MONTHS_RU[d.month]} {d.year}",
                "time": d.strftime("%H:%M"),
                "location": e.location,
            }
        )
    return jsonify({"games": rows})

#Cписок администраторов в блоке Администраторы
@api_admin_bp.route("/admin/users/admins", methods=["GET"])
def admin_list_admins():
    if not _require_admin_json():
        return jsonify({"status": "error", "message": "Нет доступа"}), 403

    admins = User.query.filter(User.role == "admin").order_by(User.email).all()
    return jsonify([
        {
            "id": u.id,
            "name": u.name,
            "email": u.email,
            "created_at": _format_date_ru(u.created_at),
        }
        for u in admins
    ])

#Очистка истории мероприятий
@api_admin_bp.route("/admin/history/clear", methods=["POST"])
def admin_clear_history():
    if not _require_admin_json():
        return jsonify({"status": "error", "message": "Нет доступа"}), 403

    from datetime import timedelta
    six_months_ago = datetime.now() - timedelta(days=180)
    past_events = Event.query.filter(Event.date < datetime.now(), Event.date >= six_months_ago).all()

    count = len(past_events)
    for e in past_events:
        if e.photo:
            photo_path = os.path.join(current_app.root_path, "media", e.photo)
            if os.path.exists(photo_path):
                os.remove(photo_path)
        if e.result_photo:
            result_photo_path = os.path.join(current_app.root_path, "media", e.result_photo)
            if os.path.exists(result_photo_path):
                os.remove(result_photo_path)
        db.session.delete(e)

    try:
        db.session.commit()
        return jsonify({"status": "success", "message": f"Удалено {count} мероприятий"})
    except Exception:
        db.session.rollback()
        return jsonify({"status": "error", "message": "Не удалось очистить историю"}), 500



# История прошедших мероприятий (последние 6 месяцев)
@api_admin_bp.route("/admin/history")
def admin_history():
    if not _require_admin_json():
        return jsonify({"status": "error", "message": "Нет доступа"}), 403

    from datetime import timedelta
    six_months_ago = datetime.now() - timedelta(days=180)
    past_events = Event.query.filter(Event.date < datetime.now(), Event.date >= six_months_ago).order_by(Event.date.desc()).all()

    result = []
    for e in past_events:
        regs = RegistrationsEvent.query.filter_by(event_id=e.id).all()
        teams_data = []
        total_players = 0
        for reg in regs:
            team = reg.team
            captains_name = team.members[0].name if team.members else ""
            teams_data.append({
                "team_name": team.name,
                "captain": captains_name,
                "player_count": reg.player_count,
                "status": reg.status,
            })
            total_players += reg.player_count

        result.append({
            "id": e.id,
            "title": e.name,
            "date": e.date.strftime("%Y-%m-%d"),
            "time": e.date.strftime("%H:%M"),
            "location": e.location,
            "category": e.category,
            "teams": teams_data,
            "total_teams": len(teams_data),
            "total_players": total_players,
            "has_scores": bool(e.scores or e.result_photo),
        })

    return jsonify(result)


# выдать / снять статус администратора по email
@api_admin_bp.route("/admin/users/role", methods=["POST"])
def admin_set_user_role():
    if not _require_admin_json():
        return jsonify({"status": "error", "message": "Нет доступа"}), 403

    data = request.get_json(silent=True) or {}
    email = (data.get("email") or "").strip().lower()
    role = (data.get("role") or "").strip().lower()

    if not email:
        return jsonify({"status": "error", "message": "Укажите email пользователя"}), 400
    if role not in ("admin", "user"):
        return jsonify({"status": "error", "message": "Недопустимая роль"}), 400

    user = User.query.filter_by(email=email).first()
    if not user:
        return jsonify({"status": "error", "message": "Пользователь с таким email не найден"}), 404

    if user.role == role:
        already = "уже является администратором" if role == "admin" else "уже не является администратором"
        return jsonify({"status": "success", "message": f"Пользователь {user.name} ({user.email}) {already}"})

    user.role = role
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        return jsonify({"status": "error", "message": "Не удалось обновить роль"}), 500

    action = "выдан статус администратора" if role == "admin" else "снят статус администратора"
    return jsonify({"status": "success", "message": f"Пользователю {user.name} ({user.email}) {action}"})



#После нужно сделать рефакторинг этого места
# Экспорт мероприятия в Excel
@api_admin_bp.route("/admin/events/<int:event_id>/export")
def admin_export_event_excel(event_id):
    if not _require_admin_json():
        return jsonify({"status": "error", "message": "Нет доступа"}), 403

    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from flask import send_file
    import tempfile

    event = Event.query.get(event_id)
    if not event:
        return jsonify({"status": "error", "message": "Мероприятие не найдено"}), 404

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Мероприятие"

    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color="F5C518", end_color="F5C518", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Информация о мероприятии
    ws.cell(row=1, column=1, value="Мероприятие:").font = Font(bold=True, size=14)
    ws.cell(row=1, column=2, value=event.name).font = Font(size=14)
    ws.cell(row=2, column=1, value="Дата:").font = header_font
    ws.cell(row=2, column=2, value=event.date.strftime("%d.%m.%Y %H:%M"))
    ws.cell(row=3, column=1, value="Место:").font = header_font
    ws.cell(row=3, column=2, value=event.location)
    ws.cell(row=4, column=1, value="Категория:").font = header_font
    ws.cell(row=4, column=2, value=event.category)

    # Таблица команд
    ws.cell(row=6, column=1, value="Команды участники:").font = Font(bold=True, size=14)

    headers = ["Команда", "Капитан", "Игроков", "Статус"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=7, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    regs = RegistrationsEvent.query.filter_by(event_id=event.id).order_by(RegistrationsEvent.id).all()
    total_players = 0
    for i, reg in enumerate(regs):
        row = 8 + i
        team = reg.team
        captain = team.members[0].name if team.members else ""
        ws.cell(row=row, column=1, value=team.name).border = thin_border
        ws.cell(row=row, column=2, value=captain).border = thin_border
        ws.cell(row=row, column=3, value=reg.player_count).border = thin_border
        status_text = "Подтверждено" if reg.status == "confirmed" else "Ожидание"
        if reg.waitlist:
            status_text = "Лист ожидания"
        ws.cell(row=row, column=4, value=status_text).border = thin_border
        total_players += reg.player_count

    # Итог
    summary_row = 8 + len(regs) + 1
    ws.cell(row=summary_row, column=1, value="Итого команд:").font = header_font
    ws.cell(row=summary_row, column=2, value=len(regs))
    ws.cell(row=summary_row + 1, column=1, value="Всего участников:").font = header_font
    ws.cell(row=summary_row + 1, column=2, value=total_players)

    # Результаты (если есть)
    scores = event.get_scores()
    if scores:
        score_row = summary_row + 3
        ws.cell(row=score_row, column=1, value="Результаты игры:").font = Font(bold=True, size=14)
        score_row += 1

        score_headers = ["Команда"] + [f"Тур {t+1}" for t in range(event.rounds)] + ["Сумма", "Место"]
        for col, h in enumerate(score_headers, 1):
            cell = ws.cell(row=score_row, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        sorted_scores = sorted(scores, key=lambda s: sum(s["scores"][:event.rounds]), reverse=True)
        for i, s in enumerate(sorted_scores):
            row = score_row + 1 + i
            ws.cell(row=row, column=1, value=s["team_name"]).border = thin_border
            total = 0
            for t in range(event.rounds):
                val = s["scores"][t] if t < len(s["scores"]) else 0
                ws.cell(row=row, column=2 + t, value=val).border = thin_border
                total += val
            ws.cell(row=row, column=2 + event.rounds, value=total).border = thin_border
            ws.cell(row=row, column=2 + event.rounds, value=total).font = Font(bold=True)
            ws.cell(row=row, column=3 + event.rounds, value=i + 1).border = thin_border

    # Автоширина колонок
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_length + 4

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    wb.save(tmp.name)
    tmp.close()

    filename = f"{event.name}_{event.date.strftime('%Y%m%d')}.xlsx"
    return send_file(tmp.name, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")